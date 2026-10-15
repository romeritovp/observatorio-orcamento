"""IDEB (Inep): extrator e publicação, com planilhas sintéticas no layout da original (sem rede).

Reproduz as armadilhas da fonte: cabeçalho em três linhas com mesclagens, ano com nota colada ("20235"), código
técnico ausente ou errado, "-" e vírgula decimal, rede com nota de rodapé ("Privada (2)"), nome abreviado
("R. G. do Sul"), regiões e Brasil pelo nome.
"""

import openpyxl
import pandas as pd
import pytest

from extractors import base
from extractors.nacional.inep import ideb
from publicacao import ideb as pub

TITULOS = {"anos_iniciais": "Ensino Fundamental Regular - Anos Iniciais", "anos_finais": "Ensino Fundamental Regular - Anos Finais",
           "medio": "Ensino Médio Regular"}
SERIES = {"anos_iniciais": ("1º ao 5º ano", ["1º", "2º", "3º", "4º", "5º"], ["SI", "1", "2", "3", "4"]),
          "anos_finais": ("6º a 9º ano", ["6º", "7º", "8º", "9º"], ["1", "2", "3", "4"]),
          "medio": ("Total", ["1ª", "2ª", "3ª"], ["1", "2", "3"])}
ANOS = [2021, 2023]  # 2021 com a nota colada no cabeçalho ("20215")


def aba(wb, nome, etapa, linhas, codigo=lambda c: c, notas=("(2) Médias do SAEB 2009 e Ideb 2009 não calculados devido à perda amostral;",)):
    """linhas: [(território, rede, {(indicador, ano, serie): valor})]. codigo(c) permite estragar um código."""
    ws = wb.create_sheet(nome)
    ws["A4"] = TITULOS[etapa]
    total, nomes, sufixos = SERIES[etapa]
    cols, c = [], 3
    for ano in ANOS:
        rotulo = f"{ano}5" if ano == 2021 else str(ano)
        inicio = c
        ws.cell(7, c, f"Taxa de Aprovação - {rotulo}")
        for serie, texto, suf in [("total", total, "SI_4"), *zip([n[0] for n in nomes], nomes, sufixos)]:
            ws.cell(9, c, texto); ws.cell(10, c, codigo(f"VL_APROVACAO_{ano}_{suf}")); cols.append((c, "aprovacao", ano, serie)); c += 1
        ws.cell(9, c, "Indicador de Rendimento (P)"); ws.cell(10, c, codigo(f"VL_INDICADOR_REND_{ano}")); cols.append((c, "rendimento", ano, None)); c += 1
        ws.merge_cells(start_row=7, start_column=inicio, end_row=8, end_column=c - 1)
    for ano in ANOS:
        ws.cell(7, c, f"Nota SAEB - {ano}")
        for texto, ind, cod in [("Matemática", "saeb_matematica", "MATEMATICA"), ("Língua Portuguesa", "saeb_portugues", "PORTUGUES"),
                                ("Nota Média Padronizada (N)", "nota_padronizada", "MEDIA")]:
            ws.cell(8, c, texto); ws.cell(10, c, codigo(f"VL_NOTA_{cod}_{ano}")); cols.append((c, ind, ano, None)); c += 1
        ws.merge_cells(start_row=7, start_column=c - 3, end_row=7, end_column=c - 1)
    for ano in ANOS:
        ws.cell(7, c, f"IDEB\n{ano}\n(N x P)"); ws.cell(10, c, codigo(f"VL_OBSERVADO_{ano}")); cols.append((c, "ideb", ano, None)); c += 1
    ws.cell(7, c, "Metas do 1º ciclo do Ideb6,7 (2007-2021)"); ws.cell(9, c, 2021); ws.cell(10, c, codigo("VL_PROJECAO_2021"))
    cols.append((c, "meta", 2021, None))
    r = 11
    for territorio, rede, valores in linhas:
        ws.cell(r, 1, territorio); ws.cell(r, 2, rede)
        for col, ind, ano, serie in cols:
            ws.cell(r, col, valores.get((ind, ano, serie), 50.0 if ind == "aprovacao" else 0.9 if ind == "rendimento"
                                        else 200.0 if ind.startswith("saeb") else 5.0 if ind == "nota_padronizada" else 4.5))
        r += 1
    ws.cell(r + 1, 1, "Fonte: MEC/Inep.")
    for i, nota in enumerate(notas):
        ws.cell(r + 2 + i, 1, ("Notas: " if i == 0 else "") + nota)
    return ws


def planilhas(tmp_path, **kw):
    """Arquivo de UFs (RS e SE, com a região Sul) e arquivo do Brasil, as três etapas em cada."""
    caminhos = []
    for nome, territorios in [("ufs", [("R. G. do Sul", "Estadual"), ("R. G. do Sul", "Privada (2)"), ("Sergipe", "Estadual"), ("Sul", "Total")]),
                              ("brasil", [("Brasil", "Total"), ("Brasil", "Municipal")])]:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for etapa, sufixo in [("anos_iniciais", "AI"), ("anos_finais", "AF"), ("medio", "EM")]:
            valores = {("ideb", 2023, None): "-", ("saeb_matematica", 2021, None): "201,5"} if etapa == "medio" else {}
            aba(wb, kw.get("titulo_ai", f"{nome} ({sufixo})") if sufixo == "AI" else f"{nome} ({sufixo})", etapa,
                [(t, r, valores) for t, r in territorios], **({"codigo": kw["codigo"]} if "codigo" in kw and sufixo == "AI" and nome == "ufs" else {}))
        wb.save(tmp_path / f"{nome}.xlsx")
        caminhos.append(tmp_path / f"{nome}.xlsx")
    return caminhos


@pytest.fixture
def lido(tmp_path):
    return ideb.ler_planilhas(planilhas(tmp_path))


def test_le_pelo_cabecalho_de_texto(lido):
    df = lido["linhas"]
    assert set(df["uf"]) == {"RS", "SE", "SUL", "BR"}
    assert set(df["rede"]) == {"estadual", "privada", "total", "municipal"}
    assert set(df["ano"]) == {2021, 2023}  # "20215" é 2021 com a nota 5
    aprovacao = df[(df["indicador"] == "aprovacao") & (df["uf"] == "RS")]
    assert set(aprovacao[aprovacao["etapa"] == "anos_iniciais"]["serie"]) == {"total", "1", "2", "3", "4", "5"}
    assert set(aprovacao[aprovacao["etapa"] == "medio"]["serie"]) == {"total", "1", "2", "3"}
    assert set(df.loc[df["indicador"] == "meta", "ano"]) == {2021}
    assert df[df["indicador"] != "aprovacao"]["serie"].isna().all()


def test_ausente_vira_nulo_e_virgula_decimal(lido):
    df = lido["linhas"]
    medio = df[(df["etapa"] == "medio") & (df["uf"] == "RS") & (df["rede"] == "estadual")]
    assert medio[(medio["indicador"] == "ideb") & (medio["ano"] == 2023)]["valor"].isna().all()  # "-" → nulo, não 0
    assert medio[(medio["indicador"] == "saeb_matematica") & (medio["ano"] == 2021)]["valor"].iloc[0] == 201.5


def test_nota_no_rotulo_da_rede_vira_ressalva_com_o_texto_do_rodape(lido):
    r = [x for x in lido["ressalvas"] if x["uf"] == "RS"]
    assert {(x["etapa"], x["rede"], x["nota"]) for x in r} == {(e, "privada", "2") for e in TITULOS}
    assert r[0]["texto"].startswith("Médias do SAEB 2009") and r[0]["anos"] == [2009]


def test_codigo_ausente_registrado_divergente_e_erro_e_erro_conhecido_aceito(tmp_path):
    sem = ideb.ler_planilhas(planilhas(tmp_path, codigo=lambda c: None if c == "VL_OBSERVADO_2023" else c))
    assert {(a["indicador"], a["ano"], a["problema"]) for a in sem["codigos_ausentes"]} == {("ideb", 2023, "código ausente")}
    with pytest.raises(base.ErroValidacao, match="ano do cabeçalho 2023"):
        ideb.ler_planilhas(planilhas(tmp_path, codigo=lambda c: "VL_OBSERVADO_2021" if c == "VL_OBSERVADO_2023" else c))
    # o erro de digitação conhecido da fonte (P de 2025 rotulado como 2023 na aba de anos iniciais) não para a leitura
    conhecido = ("UF e Regiões (AI)", "VL_INDICADOR_REND_2021", "rendimento", 2023)
    ideb.CODIGOS_COM_ERRO[conhecido] = "teste"
    try:
        ok = ideb.ler_planilhas(planilhas(tmp_path, titulo_ai="UF e Regiões (AI)",
                                          codigo=lambda c: "VL_INDICADOR_REND_2021" if c == "VL_INDICADOR_REND_2023" else c))
        assert any(a["problema"].startswith("código com erro na fonte") for a in ok["codigos_ausentes"])
    finally:
        del ideb.CODIGOS_COM_ERRO[conhecido]


def test_territorio_sem_sigla_e_erro(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for etapa in TITULOS:
        aba(wb, etapa, etapa, [("Rio Grande do Sul", "Total", {})])  # por extenso não está na tabela fixa
    wb.save(tmp_path / "x.xlsx")
    with pytest.raises(base.ErroValidacao, match="território sem sigla"):
        ideb.ler_planilhas([tmp_path / "x.xlsx"])


def test_conferencias_n_x_p_controles_e_amplitude(lido):
    df = lido["linhas"].copy()
    i = df[(df["uf"] == "RS") & (df["etapa"] == "anos_finais") & (df["rede"] == "estadual") & (df["indicador"] == "ideb") & (df["ano"] == 2021)].index
    df.loc[i, "valor"] = 4.8  # N × P = 5,0 × 0,9 = 4,5: fora de ±0,1
    r = ideb.conferencias(df, ["RS", "SE"], [2021, 2023])
    assert [(f["uf"], f["etapa"], f["ano"]) for f in r["n_x_p"]["fora_da_tolerancia"]] == [("RS", "anos_finais", 2021)]
    assert not any(c["ok"] for c in r["controles_2025"])  # sem 2025 na planilha de teste
    assert {(l["uf"], l["etapa"], l["rede"]) for l in r["cobertura"]["lacunas"]} >= {("RS", "medio", "estadual")}  # 2023 "-"
    assert all(not v for v in r["amplitude"].values())


def test_publicacao_formato_metas_e_ressalvas(lido):
    df = ideb.tabela(lido["linhas"], "2026-10-14T00:00:00+00:00", 2025)
    df = base.validar_esquema(df, "ideb")
    fonte = {"edicao": 2025, "anos": [2021, 2023], "nota_original": "Fonte: MEC/Inep.", "ressalvas_por_linha": lido["ressalvas"],
             "codigos_ausentes": [], "conferencias": {"cobertura": {}},
             "arquivos": [{"extraido_em": "2026-10-14T00:00:00+00:00", "url": "u", "sha256": "s"}]}
    ufs = pd.DataFrame({"uf": ["RS", "SE"], "nome": ["Rio Grande do Sul", "Sergipe"]})
    agregados = pd.DataFrame({"uf": ["SUL", "BR"], "nome": ["Sul", "Brasil"]})
    j = pub.montar_ideb_json(df, fonte, ufs, agregados)
    assert j["fonte"]["formulacao"] == "original_2007"
    assert list(j["agregados"]) == ["sul", "brasil"] and "municipal" in j["agregados"]["brasil"]["etapas"]["anos_iniciais"]
    rs = j["ufs"]["RS"]["etapas"]["anos_iniciais"]["estadual"]
    assert list(rs["meta"]) == ["2021"] and list(rs["aprovacao"]) == ["total", "1", "2", "3", "4", "5"]
    assert j["ufs"]["RS"]["etapas"]["medio"]["estadual"]["ideb"]["2023"] is None
    ids = {r["id"] for r in j["ressalvas"]}
    assert {"ideb_pandemia_2021", "ideb_metas_ate_2021", "ideb_medio_saeb_censitario", "ideb_formulacao_2007"} <= ids
    assert next(r for r in j["ressalvas"] if r["id"] == "ideb_pandemia_2021")["anos"] == [2021]
    assert "DF" in j["ressalvas_por_uf"] and j["ressalvas_por_uf"]["RS"][0]["linhas"][0]["rede"] == "privada"


def test_site_sem_metas_listas_alinhadas_aos_anos_e_aprovacao_por_uf(lido):
    df = base.validar_esquema(ideb.tabela(lido["linhas"], "2026-10-14T00:00:00+00:00", 2025), "ideb")
    fonte = {"edicao": 2025, "anos": [2021, 2023], "nota_original": None, "ressalvas_por_linha": lido["ressalvas"], "codigos_ausentes": [],
             "conferencias": {"cobertura": {}, "controles_2025": [], "n_x_p": {"conferidos": 1, "ideb_sem_n_ou_p": []}},
             "arquivos": [{"extraido_em": "2026-10-14T00:00:00+00:00", "url": "u", "sha256": "s"}]}
    j = pub.montar_ideb_json(df, fonte, pd.DataFrame({"uf": ["RS", "SE"], "nome": ["Rio Grande do Sul", "Sergipe"]}),
                             pd.DataFrame({"uf": ["SUL", "BR"], "nome": ["Sul", "Brasil"]}))
    site = pub.montar_ideb_site(j)
    assert set(site) == {"ideb.json", "ideb/RS.json", "ideb/SE.json"}
    geral = site["ideb.json"]
    assert "meta" not in str(geral["ufs"]) and all(i["id"] != "meta" for i in geral["indicadores"])
    rs = geral["ufs"]["RS"]["etapas"]["medio"]["estadual"]
    assert rs["ideb"] == [4.5, None] and list(rs["aprovacao"]) == ["total"]  # 2023 "-" → nulo; só a aprovação total
    assert list(site["ideb/RS.json"]["anos_iniciais"]["estadual"]["aprovacao"]) == ["total", "1", "2", "3", "4", "5"]
    meta = pub.bloco_meta(j)
    assert meta["etapa_padrao"] == "medio" and meta["rede_padrao"] == "estadual" and "municipal" not in meta["redes"]


def test_gabarito_posicao_e_variacao(lido):
    df = ideb.tabela(lido["linhas"], "x", 2025)
    i = df[(df["uf"] == "SE") & (df["etapa"] == "anos_iniciais") & (df["rede"] == "estadual") & (df["indicador"] == "ideb") & (df["ano"] == 2023)].index
    df.loc[i, "valor"] = 5.0
    g = {(x["uf"], x["etapa"], x["rede"], x["ano"]): x for x in pub.gabarito_ideb(df, ["RS", "SE"])}
    assert g[("SE", "anos_iniciais", "estadual", 2023)]["posicao"] == 1 and g[("RS", "anos_iniciais", "estadual", 2023)]["posicao"] == 2
    assert g[("SE", "anos_iniciais", "estadual", 2023)]["variacao_pontos"] == pytest.approx(0.5)
    assert g[("RS", "anos_iniciais", "estadual", 2021)]["variacao_pontos"] is None
    assert ("SUL", "anos_iniciais", "estadual", 2023) not in g  # só as UFs
