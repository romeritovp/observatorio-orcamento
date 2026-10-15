"""IBID (INPI): extrator e publicação, com uma planilha sintética no layout da original (sem rede).

A planilha de teste reproduz as armadilhas da fonte: "SE" para Sergipe (UF) e para Sudeste (região), e as
colunas N:O ("PIB pc R$10k") preenchidas, que não podem ser lidas.
"""

import openpyxl
import pandas as pd
import pytest

from extractors import base
from extractors.nacional.inpi import ibid
from publicacao import ibid as pub

UFS = pd.read_csv(base.CAMINHO_UFS, dtype=str)
SIGLAS = UFS["uf"].tolist()
DIMENSOES = pd.read_csv(ibid.CAMINHO_DIMENSOES, dtype=str, keep_default_na=False)
CABECALHO = [c or None for c in DIMENSOES["cabecalho_planilha"]]
REGIOES = [("Norte", "NO"), ("Nordeste", "NE"), ("Sudeste", "SE"), ("Sul", "SU"), ("Centro-Oeste", "CO"), ("Brasil", "BR")]
ANOS = [2024, 2025]


def escore(i_uf, i_dim, ano):
    """Escores distintos e determinísticos: a ordem das UFs é a da lista (a primeira é a melhor)."""
    return round(0.9 - i_uf * 0.03 - i_dim * 0.001 + (ano - 2024) * 0.0001, 6)


def planilha(caminho, cabecalho=CABECALHO, sem_uf=None, ranking_trocado=False):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for dim, aba in ibid.ABAS_RANKING.items():
        ws = wb.create_sheet(aba)
        ws.cell(7, 1, "UF")
        for j, ano in enumerate(ANOS):
            ws.cell(7, 2 + j, ano)
        for i, uf in enumerate(SIGLAS):
            ws.cell(9 + i, 1, uf)
            for j, ano in enumerate(ANOS):
                posicao = i + 1
                if ranking_trocado and dim == "geral" and ano == 2025 and i in (0, 1):
                    posicao = 2 - i  # troca 1º e 2º: diverge da ordenação dos escores
                ws.cell(9 + i, 2 + j, posicao)
        if dim == "geral":
            ws.cell(40, 1, "Fonte: INPI, Coordenação-Geral de Economia e Inovação.")
    for ano in ANOS:
        ws = wb.create_sheet(f"IBID {ano}")
        ws["A1"] = ano
        for k, c in enumerate(cabecalho):
            ws.cell(3, 3 + k, c)
        linha = 4
        for i, (uf, nome) in enumerate(zip(UFS["uf"], UFS["nome"])):
            if uf == sem_uf:
                continue
            ws.cell(linha, 1, nome)
            ws.cell(linha, 2, uf)
            for k in range(len(CABECALHO)):
                ws.cell(linha, 3 + k, escore(i, k, ano))
            ws.cell(linha, 14, 9.99)  # N:O: "PIB pc (R$10k)", não deve ser lido
            ws.cell(linha, 15, 1.23)
            linha += 1
        linha += 2
        for r, (nome, sigla) in enumerate(REGIOES):
            ws.cell(linha + r, 1, nome)
            ws.cell(linha + r, 2, sigla)
            for k in range(len(CABECALHO)):
                ws.cell(linha + r, 3 + k, 0.3 + r * 0.01 + k * 0.001)
    wb.save(caminho)
    return caminho


@pytest.fixture
def lido(tmp_path):
    return ibid.ler_planilha(planilha(tmp_path / "ibid.xlsx"))


def test_le_ufs_regioes_e_brasil_separando_sergipe_de_sudeste(lido):
    assert lido["anos"] == ANOS
    assert set(lido["ufs"]) == set(SIGLAS)
    assert set(lido["agregados"]) == {"NORTE", "NORDESTE", "SUDESTE", "SUL", "CENTRO_OESTE", "BR"}
    assert lido["ufs"]["SE"]["nome"] == "Sergipe"
    assert lido["agregados"]["SUDESTE"]["nome"] == "Sudeste"
    i_se = SIGLAS.index("SE")
    assert lido["ufs"]["SE"]["escores"]["geral"]["2025"] == escore(i_se, 0, 2025)
    assert lido["agregados"]["SUDESTE"]["escores"]["geral"]["2025"] == pytest.approx(0.32)


def test_dez_dimensoes_e_nada_das_colunas_n_o(lido):
    escores = lido["ufs"]["SP"]["escores"]
    assert list(escores) == DIMENSOES["dimensao"].tolist()
    assert all(v not in (9.99, 1.23) for d in escores.values() for v in d.values())


def test_ranking_oficial_conferido_com_os_escores(lido, tmp_path):
    assert lido["ufs"][SIGLAS[0]]["ranking"]["geral"]["2025"] == 1
    assert lido["divergencias"] == []
    assert lido["nota_original"] == "Fonte: INPI, Coordenação-Geral de Economia e Inovação."
    trocado = ibid.ler_planilha(planilha(tmp_path / "trocado.xlsx", ranking_trocado=True))
    assert {(d["uf"], d["oficial"], d["pelo_escore"]) for d in trocado["divergencias"]} == {
        (SIGLAS[0], 2, 1), (SIGLAS[1], 1, 2)}


@pytest.mark.parametrize("estrago, mensagem", [
    ({"cabecalho": [None, "IBID - Contexto (novo)", *CABECALHO[2:]]}, "cabeçalho mudou"),
    ({"sem_uf": "RJ"}, "faltam UFs"),
])
def test_layout_diferente_falha(tmp_path, estrago, mensagem):
    with pytest.raises(base.ErroValidacao, match=mensagem):
        ibid.ler_planilha(planilha(tmp_path / "x.xlsx", **estrago))


def test_tabelas_no_esquema_com_codigos_dos_dicionarios(lido):
    escore, ranking = ibid.tabelas(lido, "2026-10-14T00:00:00+00:00")
    escore, ranking = base.validar_esquema(escore, "ibid_escore"), base.validar_esquema(ranking, "ibid_ranking")
    assert len(escore) == (27 + 6) * len(DIMENSOES) * len(ANOS)
    assert len(ranking) == 27 * 3 * len(ANOS)
    codigos = dict(zip(escore["uf"], escore["cod_ibge"]))
    assert codigos["SE"] == "28" and codigos["SUDESTE"] == "3" and codigos["BR"] == "1"
    assert set(escore["fonte_dado"]) == {"inpi_ibid"}


def test_publicacao_mesmo_formato_da_saida_original(lido):
    tabela_escore, tabela_ranking = ibid.tabelas(lido, "2026-10-14T15:10:30+00:00")
    agregados = pd.concat([pd.read_csv(ibid.CAMINHO_REGIOES, dtype=str), pd.read_csv(ibid.CAMINHO_TERRITORIOS, dtype=str)])
    fonte = {"url": "u", "arquivo_sha256": "abc", "extraido_em": "2026-10-14T15:10:30+00:00",
             "nota_original": lido["nota_original"], "divergencias_ranking": []}
    config = {"nome": "IBID", "orgao": "INPI", "pagina_oficial": "p", "escores_comparaveis_entre_anos": False}
    j = pub.montar_ibid_json(tabela_escore, tabela_ranking, DIMENSOES, UFS, agregados, fonte, config)
    assert list(j["ufs"]) == sorted(SIGLAS)
    assert list(j["agregados"]) == ["norte", "nordeste", "sudeste", "sul", "centro_oeste", "brasil"]
    assert j["fonte"]["extraido_em"] == "2026-10-14"
    assert j["dimensoes"][0] == {"id": "geral", "nome": "IBID – Geral", "pai": None}
    assert j["ufs"]["SE"]["ranking"]["geral"] == {"2024": SIGLAS.index("SE") + 1, "2025": SIGLAS.index("SE") + 1}
    assert j["conferencia"] == {"ranking_oficial_vs_escore": "ok", "divergencias": []}
    # Estado típico: mediana das 27 UFs, calculada aqui (não pelo INPI), por dimensão e ano
    tipico = j["mediana_ufs"]
    assert tipico["nome"] == "Estado típico (mediana das UFs)" and "não é um número do INPI" in tipico["calculado_por"]
    i_meio = sorted(range(27), key=lambda i: -escore(i, 0, 2025))[13]  # 14º de 27
    assert tipico["escores"]["geral"]["2025"] == pytest.approx(escore(i_meio, 0, 2025))
    assert set(tipico["escores"]) == set(DIMENSOES["dimensao"])
    # Comparabilidade dos escores: um único ponto de configuração decide a ressalva
    ids = lambda c: [r["id"] for r in pub.bloco_meta(j, {**config, "escores_comparaveis_entre_anos": c})["ressalvas"]]
    assert "ibid_escores_entre_anos" in ids(False) and "ibid_escores_entre_anos" not in ids(True)
    assert pub.bloco_meta(j, config)["escores_comparaveis_entre_anos"] is False


# --- códigos das regiões: a tabela do IBID se liga pela coluna uf, nunca por cod_ibge ---------------------

def test_cod_ibge_do_ibid_e_ambiguo_e_ninguem_usa():
    """O Norte (região 1 do IBGE) e o Brasil (1) dividem o cod_ibge na tabela do IBID. O Brasil continua com 1 porque
    é o código usado no pib_uf e o id_ente da União no SICONFI. Este teste falha se o código do IBID, a publicação
    ou alguma tela passar a usar cod_ibge para ligar territórios do IBID (a ligação é pela coluna uf)."""
    from pathlib import Path

    raiz = base.RAIZ
    regioes = pd.read_csv(ibid.CAMINHO_REGIOES, dtype=str)
    territorios = pd.read_csv(ibid.CAMINHO_TERRITORIOS, dtype=str)
    codigos = pd.concat([regioes, territorios])["cod_ibge"]
    assert codigos.duplicated().any(), "se os códigos deixarem de colidir, revise este teste e docs/dicionario.md"
    assert regioes["uf"].str.len().gt(2).all(), "região com uf de duas letras colide com sigla de UF (ex.: SE)"

    trechos = {
        "publicacao/ibid.py": (raiz / "publicacao" / "ibid.py").read_text(encoding="utf-8"),
        "web/src/telas/PesquisaInovacao.jsx": (raiz / "web" / "src" / "telas" / "PesquisaInovacao.jsx").read_text(encoding="utf-8"),
    }
    calculos = (raiz / "web" / "src" / "calculos.js").read_text(encoding="utf-8")
    inicio = calculos.index("// --- Pesquisa e Inovação")
    fim = calculos.find("\n// --- ", inicio + 1)
    trechos["web/src/calculos.js (Pesquisa e Inovação)"] = calculos[inicio: None if fim < 0 else fim]
    for caminho in (raiz / "web" / "src" / "componentes").glob("*.jsx"):
        texto = caminho.read_text(encoding="utf-8")
        if "PesquisaInovacao" in texto or caminho.stem in ("Dispersao", "BarrasComparadas", "TabelaUfs"):
            trechos[str(caminho.relative_to(raiz))] = texto
    for nome, texto in trechos.items():
        assert "cod_ibge" not in texto, f"{nome} usa cod_ibge; ligue territórios do IBID pela coluna uf"

    # O que o front recebe não tem cod_ibge: não há como uma tela usá-lo
    publicado = raiz / "web" / "public" / "data" / "ibid.json"
    if publicado.exists():
        assert "cod_ibge" not in publicado.read_text(encoding="utf-8")
