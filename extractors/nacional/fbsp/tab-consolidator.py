"""
Junta uma aba (ex.: T02) de todos os arquivos anuario-YYYY.xlsx de uma pasta
em um único arquivo <ABA>-<primeiro_ano>-<ultimo_ano>.xlsx,
com uma aba por ano.

Uso:
    python tab-consolidator.py --tab T02                    # usa a pasta atual
    python tab-consolidator.py --tab T02,T2 /caminho/pasta  # aceita variações
    python tab-consolidator.py --tab T02 --tab T2 /caminho/pasta

O arquivo de saída usa o primeiro nome da lista: T02-<ano_ini>-<ano_fim>.xlsx

Requer: pip install openpyxl
"""

import argparse
import re
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook

PADRAO = re.compile(r"^anuario-(\d{4})\.xlsx$", re.IGNORECASE)


def copiar_aba(origem, destino):
    """Copia valores, estilos, mesclagens, larguras e alturas."""
    for linha in origem.iter_rows():
        for cel in linha:
            nova = destino.cell(row=cel.row, column=cel.column, value=cel.value)
            if cel.has_style:
                nova.font = copy(cel.font)
                nova.border = copy(cel.border)
                nova.fill = copy(cel.fill)
                nova.number_format = cel.number_format
                nova.protection = copy(cel.protection)
                nova.alignment = copy(cel.alignment)

    for intervalo in origem.merged_cells.ranges:
        destino.merge_cells(str(intervalo))

    for letra, dim in origem.column_dimensions.items():
        destino.column_dimensions[letra].width = dim.width
    for idx, dim in origem.row_dimensions.items():
        destino.row_dimensions[idx].height = dim.height

    destino.freeze_panes = origem.freeze_panes


def main():
    parser = argparse.ArgumentParser(
        description="Junta uma aba de vários anuario-YYYY.xlsx em um único arquivo."
    )
    parser.add_argument("--tab", "--tap", required=True, dest="tab", action="append",
                        help="nome(s) da aba, separados por vírgula ou repetindo "
                             "a opção. Ex.: --tab T02,T2  ou  --tab T02 --tab T2. "
                             "Em cada arquivo vale o primeiro nome encontrado.")
    parser.add_argument("pasta", nargs="?", default=".",
                        help="pasta com os arquivos (padrão: pasta atual)")
    args = parser.parse_args()

    nomes = [n.strip() for item in args.tab for n in item.split(",") if n.strip()]
    pasta = Path(args.pasta)

    if not pasta.is_dir():
        print(f"Pasta não encontrada: {pasta.resolve()}")
        return

    arquivos = []
    for f in pasta.iterdir():
        m = PADRAO.match(f.name)
        if m:
            arquivos.append((int(m.group(1)), f))
    arquivos.sort()

    if not arquivos:
        print(f"Nenhum arquivo anuario-YYYY.xlsx encontrado em {pasta}")
        return

    saida_wb = Workbook()
    saida_wb.remove(saida_wb.active)  # remove a aba vazia padrão

    anos_ok = []
    for ano, caminho in arquivos:
        # data_only=True traz o valor calculado das fórmulas,
        # evitando referências quebradas a outras abas do arquivo original
        wb = load_workbook(caminho, data_only=True)
        # compara sem diferenciar maiúsculas e ignorando espaços nas pontas
        abas = {s.strip().upper(): s for s in wb.sheetnames}
        encontrada = next((abas[n.upper()] for n in nomes if n.upper() in abas), None)
        if encontrada is None:
            print(f"[aviso] {caminho.name}: nenhuma aba {nomes}, ignorado")
            continue
        copiar_aba(wb[encontrada], saida_wb.create_sheet(title=str(ano)))
        anos_ok.append(ano)
        print(f"[ok] {caminho.name} [{encontrada}] -> aba {ano}")

    if not anos_ok:
        print("Nada foi copiado.")
        return

    nome_saida = pasta / f"{nomes[0]}-{anos_ok[0]}-{anos_ok[-1]}.xlsx"
    saida_wb.save(nome_saida)
    print(f"\nSalvo: {nome_saida} ({len(anos_ok)} abas)")


if __name__ == "__main__":
    main()