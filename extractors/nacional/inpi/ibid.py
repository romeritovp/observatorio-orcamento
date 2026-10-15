"""Extrator do IBID (Índice Brasil de Inovação e Desenvolvimento) — INPI.

Fonte: https://www.gov.br/inpi/pt-br/inpi-data/indice-brasil-de-inovacao-e-desenvolvimento-ibid/ibid-serie-historica.xlsx

Lê as abas anuais ("IBID 2014" … "IBID 2025") com os escores (0 a 1) do índice
geral, dos dois subíndices e dos sete pilares, por UF e por agregado (regiões e
Brasil), e as três abas de série histórica com a posição oficial no ranking.

Uso:
    uv run python -m extractors.nacional.inpi.ibid              # baixa a planilha se faltar e grava
    uv run python -m extractors.nacional.inpi.ibid --atualizar  # baixa de novo

Bruto: data/raw/inpi/ibid/ibid-serie-historica.xlsx (extractors.nacional.inpi.baixar, com manifesto).
Tabelas gravadas: `ibid_escore` (27 UFs, 5 regiões e o Brasil) e `ibid_ranking` (27 UFs), com escores sem
arredondar; e data/processed/ibid_fonte.json (nota da planilha, SHA-256 do arquivo, data da extração e a
conferência do ranking). O JSON do site é montado por publicacao.ibid.

Cuidados da fonte:
- A sigla "SE" aparece duas vezes em cada aba anual: Sergipe (linha de UF) e
  Sudeste (linha de região). Por isso os agregados são lidos pelo nome, nunca
  pela sigla.
- As colunas N:O de cada aba anual trazem um "PIB pc (R$10k)" auxiliar, sem
  fonte nem unidade clara. Não é extraído: o PIB do projeto vem do IBGE.
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
from pathlib import Path

import openpyxl
import pandas as pd

from extractors import base
from extractors.nacional.inpi import baixar as baixar_mod

FONTE_DADO = "inpi_ibid"
CAMINHO_DIMENSOES = base.RAIZ / "dicionarios" / "nacional" / "ibid_dimensoes.csv"
CAMINHO_REGIOES = base.RAIZ / "dicionarios" / "nacional" / "regioes.csv"
CAMINHO_TERRITORIOS = base.RAIZ / "dicionarios" / "nacional" / "territorios.csv"
CAMINHO_FONTE = base.DIR_PROCESSED / "ibid_fonte.json"

ABAS_RANKING = {
    "geral": "IBID_série hist",
    "contexto": "IBID-contexto_série hist",
    "resultado": "IBID-resultado_série hist",
}


def conferir(condicao, mensagem):
    """As conferências do layout: falham sempre (um `assert` some com `python -O`)."""
    if not condicao:
        raise base.ErroValidacao(mensagem)


def _limpa(v):
    return v.strip() if isinstance(v, str) else v


def dicionarios(caminho_ufs=base.CAMINHO_UFS, caminho_dimensoes=CAMINHO_DIMENSOES,
                caminho_regioes=CAMINHO_REGIOES, caminho_territorios=CAMINHO_TERRITORIOS):
    """(siglas das 27 UFs, dimensões [(id, cabeçalho da planilha)] na ordem das colunas C..L,
    {nome do agregado na planilha: uf do agregado}) a partir dos dicionários versionados."""
    ufs = pd.read_csv(caminho_ufs, dtype=str)["uf"].tolist()
    dims = pd.read_csv(caminho_dimensoes, dtype=str, keep_default_na=False)
    dimensoes = [(d, c or None) for d, c in zip(dims["dimensao"], dims["cabecalho_planilha"])]
    agregados = pd.concat([pd.read_csv(caminho_regioes, dtype=str), pd.read_csv(caminho_territorios, dtype=str)])
    return ufs, dimensoes, dict(zip(agregados["nome"], agregados["uf"]))


def ler_anuais(wb, ufs_validas, dimensoes, nomes_agregados):
    """Escores das abas anuais: (anos, {uf: {"nome", "escores": {dim: {ano: valor}}}}, {agregado: idem})."""
    anos = sorted(
        int(m.group(1)) for n in wb.sheetnames if (m := re.fullmatch(r"IBID (\d{4})", n))
    )
    conferir(anos, "nenhuma aba anual 'IBID AAAA'")
    cabecalho_esperado = [c for _, c in dimensoes]
    ufs: dict[str, dict] = {}
    agregados: dict[str, dict] = {}
    for ano in anos:
        ws = wb[f"IBID {ano}"]
        conferir(ws["A1"].value == ano, f"aba {ano}: A1={ws['A1'].value!r}")
        cab = [_limpa(c.value) for c in ws[3][2:2 + len(dimensoes)]]
        conferir(cab == cabecalho_esperado, f"aba {ano}: cabeçalho mudou: {cab}")

        vistos = set()
        for row in ws.iter_rows(min_row=4, max_row=ws.max_row, max_col=2 + len(dimensoes), values_only=True):
            nome, sigla, *valores = (_limpa(v) for v in row)
            if nome is None:
                continue
            if nome in nomes_agregados:  # pelo nome: "SE" é Sergipe e também Sudeste
                destino = agregados.setdefault(nomes_agregados[nome], {"nome": nome, "escores": {}})
            elif sigla in ufs_validas and sigla not in vistos:
                vistos.add(sigla)
                destino = ufs.setdefault(sigla, {"nome": nome, "escores": {}, "ranking": {}})
                conferir(destino["nome"] == nome, f"{sigla}: nome mudou ({destino['nome']} → {nome})")
            else:
                raise base.ErroValidacao(f"aba {ano}: linha inesperada {nome!r}/{sigla!r}")
            for (dim, _), v in zip(dimensoes, valores):
                conferir(isinstance(v, (int, float)), f"aba {ano} {nome} {dim}: {v!r}")
                destino["escores"].setdefault(dim, {})[str(ano)] = float(v)
        faltando = set(ufs_validas) - vistos
        conferir(not faltando, f"aba {ano}: faltam UFs {faltando}")
    conferir(set(nomes_agregados.values()) == set(agregados), "faltam agregados")
    return anos, ufs, agregados


def ler_rankings(wb, anos, ufs, ufs_validas):
    """Posição oficial (1 = melhor) das abas de série histórica, gravada em ufs[uf]["ranking"][dim][ano]."""
    for dim, aba in ABAS_RANKING.items():
        ws = wb[aba]
        cab = [c.value for c in ws[7]]
        conferir(cab[0] == "UF" and cab[1 : 1 + len(anos)] == anos, f"{aba}: cabeçalho {cab}")
        lidas = set()
        for row in ws.iter_rows(min_row=8, values_only=True):
            sigla = _limpa(row[0])
            if sigla not in ufs_validas:
                continue
            lidas.add(sigla)
            ufs[sigla]["ranking"][dim] = {
                str(a): int(p) for a, p in zip(anos, row[1 : 1 + len(anos)])
            }
        conferir(lidas == set(ufs_validas), f"{aba}: faltam {set(ufs_validas) - lidas}")


def conferir_rankings(anos, ufs):
    """O ranking oficial bate com a ordenação dos escores? Retorna divergências."""
    divergencias = []
    for dim in ABAS_RANKING:
        for ano in map(str, anos):
            ordem = sorted(ufs, key=lambda u: -ufs[u]["escores"][dim][ano])
            for pos, uf in enumerate(ordem, 1):
                oficial = ufs[uf]["ranking"][dim][ano]
                if oficial != pos:
                    divergencias.append({"dimensao": dim, "ano": int(ano), "uf": uf,
                                         "oficial": oficial, "pelo_escore": pos})
    return divergencias


def nota_da_fonte(wb):
    """Linha "Fonte: ..." da aba da série histórica do índice geral (a última, se houver mais de uma)."""
    nota = None
    for row in wb["IBID_série hist"].iter_rows(values_only=True):
        if isinstance(row[0], str) and row[0].startswith("Fonte:"):
            nota = row[0].strip()
    return nota


def ler_planilha(caminho, dicionarios_=None):
    """Lê e confere a planilha. Devolve {anos, ufs, agregados, divergencias, nota_original}."""
    ufs_validas, dimensoes, nomes_agregados = dicionarios_ or dicionarios()
    wb = openpyxl.load_workbook(caminho, data_only=True, read_only=False)
    anos, ufs, agregados = ler_anuais(wb, ufs_validas, dimensoes, nomes_agregados)
    ler_rankings(wb, anos, ufs, ufs_validas)
    return {"anos": anos, "ufs": ufs, "agregados": agregados, "divergencias": conferir_rankings(anos, ufs),
            "nota_original": nota_da_fonte(wb)}


def tabelas(lido, extraido_em, caminho_ufs=base.CAMINHO_UFS, caminho_regioes=CAMINHO_REGIOES,
            caminho_territorios=CAMINHO_TERRITORIOS):
    """(ibid_escore, ibid_ranking) em formato longo, com uf e cod_ibge dos dicionários."""
    codigos = pd.concat([pd.read_csv(c, dtype=str)[["uf", "cod_ibge"]]
                         for c in (caminho_ufs, caminho_regioes, caminho_territorios)])
    codigo = dict(zip(codigos["uf"], codigos["cod_ibge"]))
    escores = [{"uf": uf, "ano": int(ano), "dimensao": dim, "escore": v}
               for grupo in (lido["ufs"], lido["agregados"]) for uf, d in grupo.items()
               for dim, por_ano in d["escores"].items() for ano, v in por_ano.items()]
    ranking = [{"uf": uf, "ano": int(ano), "dimensao": dim, "posicao": p}
               for uf, d in lido["ufs"].items() for dim, por_ano in d["ranking"].items() for ano, p in por_ano.items()]
    saida = []
    for linhas in (escores, ranking):
        df = pd.DataFrame(linhas)
        df["cod_ibge"] = df["uf"].map(codigo)
        conferir(df["cod_ibge"].notna().all(), f"sem código: {sorted(df.loc[df['cod_ibge'].isna(), 'uf'].unique())}")
        saida.append(df.assign(fonte_dado=FONTE_DADO, extraido_em=extraido_em))
    return tuple(saida)


def main(argv=None):
    argv = sys.argv[1:] if argv is None else argv
    caminho = baixar_mod.baixar(atualizar="--atualizar" in argv)
    registro = json.loads((baixar_mod.DIR_RAW / baixar_mod.MANIFESTO).read_text(encoding="utf-8"))[baixar_mod.RELATIVO]
    sha256 = hashlib.sha256(caminho.read_bytes()).hexdigest()
    conferir(sha256 == registro["sha256"], f"{caminho.name}: SHA-256 diferente do manifesto")
    lido = ler_planilha(caminho)
    escore, ranking = tabelas(lido, registro["extraido_em"])
    gravados = base.gravar(escore, "ibid_escore", FONTE_DADO) + base.gravar(ranking, "ibid_ranking", FONTE_DADO)
    CAMINHO_FONTE.write_text(json.dumps({
        "url": registro["url"], "arquivo": baixar_mod.RELATIVO, "arquivo_sha256": sha256,
        "extraido_em": registro["extraido_em"], "nota_original": lido["nota_original"],
        "divergencias_ranking": lido["divergencias"],
    }, ensure_ascii=False, indent=1), encoding="utf-8")
    anos = lido["anos"]
    print(f"{len(lido['ufs'])} UFs, {len(lido['agregados'])} agregados, anos {anos[0]}–{anos[-1]}; "
          f"divergências de ranking: {len(lido['divergencias'])}; {len(gravados)} arquivos gravados")


if __name__ == "__main__":
    main()
