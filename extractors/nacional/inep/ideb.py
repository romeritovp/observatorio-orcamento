"""Extrator do IDEB (Índice de Desenvolvimento da Educação Básica) — Inep, edição 2025 (série 2005–2025).

Uso:
    uv run python -m extractors.nacional.inep.ideb              # baixa os zips se faltarem, confere e grava
    uv run python -m extractors.nacional.inep.ideb --atualizar  # baixa de novo

Brutos: data/raw/inep/ideb/<edicao>/ (extractors.nacional.inep.baixar, com manifesto e SHA-256): o zip de regiões e
UFs (27 UFs e 5 regiões) e o do Brasil. Cada um traz uma planilha com uma aba por etapa (anos iniciais, anos finais,
ensino médio).
Tabela gravada: `ideb` (formato longo: território × ano × etapa × rede × indicador [× série]), valores como vieram
(aprovação em %, rendimento P de 0 a 1, notas Saeb na escala de cada etapa, nota padronizada N e IDEB de 0 a 10).
Também data/processed/ideb_fonte.json: arquivos e hashes, notas de rodapé por linha, avisos e as conferências.
O JSON de saída (data/nacional/inep/ideb.json) é montado por publicacao.ideb.

Como a planilha é lida (armadilhas em docs/fontes/inep_ideb.md):
- Colunas pelo TEXTO do cabeçalho (linhas 7 a 9, desfeitas as mesclagens), nunca pela posição. O código técnico da
  linha 10 (VL_OBSERVADO_2025...) só confere: se existe, tem de concordar; se falta, a coluna entra e a falta é
  registrada (o Brasil, anos finais, não tem o código do IDEB 2023).
- Ano do cabeçalho com nota colada ("20215" = 2021 + nota 5): lê-se os quatro primeiros dígitos.
- Território pelo NOME (dicionarios/nacional/ideb_territorios.csv; nomes abreviados como "R. G. do Sul"); nome fora
  da tabela é erro. A sigla "SE" do IBID não aparece aqui, mas regiões e Brasil também só são aceitos pelo nome.
- Rede com nota de rodapé no rótulo ("Privada (1)", "Total (3)(4)"): a nota é lida no rodapé da própria aba (a
  numeração muda entre abas e arquivos) e gravada como ressalva da linha.
- "-" e "ND" viram nulo (nunca 0); vírgula decimal é aceita; qualquer outro texto é erro.
- Metas (2007–2021) entram com indicador "meta", para não serem confundidas com o observado; 2023 e 2025 não têm meta
  oficial e nada é inventado.
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
import unicodedata

import openpyxl
import pandas as pd

from extractors import base
from extractors.nacional.inep import baixar as baixar_mod

FONTE_DADO = "inep_ideb"
CAMINHO_TERRITORIOS = base.RAIZ / "dicionarios" / "nacional" / "ideb_territorios.csv"
CAMINHO_REGIOES = base.RAIZ / "dicionarios" / "nacional" / "regioes.csv"
CAMINHO_TERRITORIOS_NAO_UF = base.RAIZ / "dicionarios" / "nacional" / "territorios.csv"
CAMINHO_FONTE = base.DIR_PROCESSED / "ideb_fonte.json"

# Etapa pelo título da aba (linha 4), conferido pelo nome da aba
ETAPAS = {
    "ensino fundamental regular - anos iniciais": "anos_iniciais",
    "ensino fundamental regular - anos finais": "anos_finais",
    "ensino medio regular": "medio",
}
REDES = {"total": "total", "estadual": "estadual", "municipal": "municipal", "publica": "publica", "privada": "privada"}
AUSENTES = {"-", "ND", "nd", "--", ""}
# Família do código técnico (linha 10) → indicador, para conferir o que o cabeçalho de texto disse
CODIGOS = [
    (re.compile(r"^VL_APROVACAO_(\d{4})_SI_4$"), "aprovacao_total"),
    (re.compile(r"^VL_APROVACAO_(\d{4})_(SI|\d)$"), "aprovacao_serie"),
    (re.compile(r"^VL_INDICADOR_REND_(\d{4})$"), "rendimento"),
    (re.compile(r"^VL_NOTA_MATEMATICA_(\d{4})$"), "saeb_matematica"),
    (re.compile(r"^VL_NOTA_PORTUGUES_(\d{4})$"), "saeb_portugues"),
    (re.compile(r"^VL_NOTA_MEDIA_(\d{4})$"), "nota_padronizada"),
    (re.compile(r"^VL_OBSERVADO_(\d{4})$"), "ideb"),
    (re.compile(r"^VL_PROJECAO_(\d{4})$"), "meta"),
]
LINHA_TITULO, LINHA_GRUPO, LINHA_SUBGRUPO, LINHA_DETALHE, LINHA_CODIGO = 4, 7, 8, 9, 10

# Erros de digitação conhecidos no código técnico (linha 10), em que vale o cabeçalho de texto. Cada um foi conferido
# à mão: a coluna está dentro do bloco mesclado do ano do texto e os valores são desse ano (o N × P confere).
# Chave: (aba, código como veio, indicador e ano pelo texto). Qualquer outra divergência é erro.
CODIGOS_COM_ERRO = {
    ("UF e Regiões (AI)", "VL_INDICADOR_REND_2023", "rendimento", 2025):
        "coluna CA, no bloco 'Taxa de Aprovação - 2025' (BU:CA); o código deveria ser VL_INDICADOR_REND_2025",
}


def conferir(condicao, mensagem):
    if not condicao:
        raise base.ErroValidacao(mensagem)


def normalizar(texto):
    """Sem acento, minúsculas, espaços simples (para comparar títulos e rótulos)."""
    sem = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", sem).strip().lower()


def territorios(caminho=CAMINHO_TERRITORIOS):
    """{nome normalizado da planilha: uf} (27 UFs, regiões por extenso e BR)."""
    df = pd.read_csv(caminho, dtype=str)
    conferir(not df["nome_planilha"].map(normalizar).duplicated().any(), "nome repetido em ideb_territorios.csv")
    return {normalizar(n): u for n, u in zip(df["nome_planilha"], df["uf"])}


def valor(celula, contexto):
    """Número da célula; "-"/"ND"/vazia → None (nunca 0); vírgula decimal aceita; outro texto é erro."""
    if celula is None:
        return None
    if isinstance(celula, bool):
        raise base.ErroValidacao(f"{contexto}: valor lógico inesperado {celula!r}")
    if isinstance(celula, (int, float)):
        return float(celula)
    texto = str(celula).strip()
    if texto in AUSENTES:
        return None
    try:
        return float(texto.replace(".", "").replace(",", ".")) if "," in texto else float(texto)
    except ValueError:
        raise base.ErroValidacao(f"{contexto}: texto inesperado {celula!r}") from None


def _preenchidas(ws, linha):
    """Valores da linha com as células mescladas preenchidas pelo valor da célula de cima/esquerda."""
    valores = {c: ws.cell(linha, c).value for c in range(1, ws.max_column + 1)}
    for faixa in ws.merged_cells.ranges:
        if faixa.min_row <= linha <= faixa.max_row:
            topo = ws.cell(faixa.min_row, faixa.min_col).value
            for c in range(faixa.min_col, faixa.max_col + 1):
                valores[c] = topo
    return valores


def _ano(texto):
    m = re.search(r"(?<!\d)(20\d{2})", str(texto))
    return int(m.group(1)) if m else None


def colunas(ws):
    """{coluna: {indicador, ano, serie, codigo}} pelo cabeçalho de texto; confere com o código técnico.
    Devolve (colunas, codigos_ausentes)."""
    grupo, sub, detalhe = (_preenchidas(ws, linha) for linha in (LINHA_GRUPO, LINHA_SUBGRUPO, LINHA_DETALHE))
    saida, ausentes = {}, []
    for c in range(3, ws.max_column + 1):
        g, s, d = (normalizar(x) if x is not None else "" for x in (grupo[c], sub[c], detalhe[c]))
        if not g:
            continue
        serie = None
        if g.startswith("taxa de aprovacao"):
            if d.startswith("indicador de rendimento"):
                indicador = "rendimento"
            else:
                indicador = "aprovacao"
                # normalizado: "1º ao 5º ano" → "1o ao 5o ano", "6º a 9º ano", "Total"; série: "3º" → "3o", "2ª" → "2a"
                serie_unica = re.fullmatch(r"(\d)[oa]?", d)
                if d == "total" or re.fullmatch(r"\d[oa]? (?:ao|a) \d[oa]? ano", d):
                    serie = "total"
                elif serie_unica:
                    serie = serie_unica.group(1)
                else:
                    raise base.ErroValidacao(f"{ws.title}, coluna {c}: série de aprovação desconhecida {detalhe[c]!r}")
            ano = _ano(grupo[c])
        elif g.startswith("nota saeb"):
            indicador = {"matematica": "saeb_matematica", "lingua portuguesa": "saeb_portugues",
                         "nota media padronizada (n)": "nota_padronizada"}.get(s)
            conferir(indicador is not None, f"{ws.title}, coluna {c}: nota Saeb desconhecida {sub[c]!r}")
            ano = _ano(grupo[c])
        elif g.startswith("ideb"):
            indicador, ano = "ideb", _ano(grupo[c])
        elif g.startswith("metas"):
            indicador, ano = "meta", _ano(detalhe[c])
        else:
            raise base.ErroValidacao(f"{ws.title}, coluna {c}: cabeçalho desconhecido {grupo[c]!r}")
        conferir(ano is not None, f"{ws.title}, coluna {c}: sem ano no cabeçalho ({grupo[c]!r} / {detalhe[c]!r})")

        codigo = ws.cell(LINHA_CODIGO, c).value
        if codigo is None:
            ausentes.append({"aba": ws.title, "coluna": c, "indicador": indicador, "ano": ano, "problema": "código ausente"})
        elif (ws.title, str(codigo).strip(), indicador, ano) in CODIGOS_COM_ERRO:
            ausentes.append({"aba": ws.title, "coluna": c, "indicador": indicador, "ano": ano, "codigo": codigo,
                             "problema": "código com erro na fonte: " + CODIGOS_COM_ERRO[(ws.title, str(codigo).strip(), indicador, ano)]})
        else:
            familia = next(((f, int(m.group(1))) for r, f in CODIGOS if (m := r.match(str(codigo).strip()))), None)
            conferir(familia is not None, f"{ws.title}, coluna {c}: código desconhecido {codigo!r}")
            esperado = {"aprovacao_total": ("aprovacao", "total")}.get(familia[0])
            if esperado:
                conferir((indicador, serie) == esperado, f"{ws.title}, coluna {c}: cabeçalho ({indicador}, {serie}) × código {codigo}")
            elif familia[0] == "aprovacao_serie":
                conferir(indicador == "aprovacao" and serie != "total", f"{ws.title}, coluna {c}: cabeçalho × código {codigo}")
            else:
                conferir(indicador == familia[0], f"{ws.title}, coluna {c}: cabeçalho ({indicador}) × código {codigo}")
            conferir(ano == familia[1], f"{ws.title}, coluna {c}: ano do cabeçalho {ano} × código {codigo}")
        saida[c] = {"indicador": indicador, "ano": ano, "serie": serie, "codigo": codigo}
    chaves = [(v["indicador"], v["ano"], v["serie"]) for v in saida.values()]
    conferir(len(chaves) == len(set(chaves)), f"{ws.title}: colunas repetidas no cabeçalho")
    return saida, ausentes


def notas_de_rodape(ws, ultima_linha):
    """{número: texto} das notas abaixo dos dados ("Notas: (1) ...", "(2) ...", com continuações)."""
    notas, atual = {}, None
    for r in range(ultima_linha + 1, ws.max_row + 1):
        texto = ws.cell(r, 1).value
        if not isinstance(texto, str) or not texto.strip():
            continue
        texto = texto.strip()
        m = re.match(r"^(?:Notas:\s*)?\((\d+)\)\s*(.*)$", texto)
        if m:
            atual = m.group(1)
            notas[atual] = m.group(2).strip()
        elif atual and not texto.startswith("Fonte"):
            notas[atual] = f"{notas[atual]} {texto}"
    return notas


def ler_aba(ws, mapa_territorios, arquivo):
    """Linhas longas da aba e as ressalvas por linha (notas no rótulo da rede)."""
    etapa = ETAPAS.get(normalizar(ws.cell(LINHA_TITULO, 1).value or ""))
    conferir(etapa is not None, f"{arquivo} / {ws.title}: título de etapa desconhecido {ws.cell(LINHA_TITULO, 1).value!r}")
    cols, ausentes = colunas(ws)
    linhas, rotulos, ultima = [], [], LINHA_CODIGO
    vistos = set()
    for r in range(LINHA_CODIGO + 1, ws.max_row + 1):
        nome, rede = ws.cell(r, 1).value, ws.cell(r, 2).value
        if nome is None or rede is None:
            continue
        ultima = r
        uf = mapa_territorios.get(normalizar(nome))
        conferir(uf is not None, f"{arquivo} / {ws.title}, linha {r}: território sem sigla {nome!r} (ideb_territorios.csv)")
        m = re.fullmatch(r"\s*([^(]+?)\s*((?:\(\d+\)\s*)*)", str(rede))
        conferir(m is not None, f"{ws.title}, linha {r}: rede ilegível {rede!r}")
        rede_id = REDES.get(normalizar(m.group(1)))
        conferir(rede_id is not None, f"{ws.title}, linha {r}: rede desconhecida {rede!r}")
        conferir((uf, rede_id) not in vistos, f"{ws.title}: {uf} {rede_id} repetido")
        vistos.add((uf, rede_id))
        for n in re.findall(r"\((\d+)\)", m.group(2)):
            rotulos.append({"uf": uf, "etapa": etapa, "rede": rede_id, "nota": n})
        for c, col in cols.items():
            linhas.append({"uf": uf, "ano": col["ano"], "etapa": etapa, "rede": rede_id, "indicador": col["indicador"],
                           "serie": col["serie"], "valor": valor(ws.cell(r, c).value, f"{ws.title} {nome} {rede} {col}")})
    notas = notas_de_rodape(ws, ultima)
    ressalvas = []
    for rotulo in rotulos:
        conferir(rotulo["nota"] in notas, f"{ws.title}: nota ({rotulo['nota']}) sem texto no rodapé")
        texto = notas[rotulo["nota"]]
        ressalvas.append({**rotulo, "texto": texto, "anos": sorted({int(a) for a in re.findall(r"(?<!\d)(20\d{2})(?!\d)", texto)})})
    return linhas, ressalvas, [{**a, "arquivo": arquivo} for a in ausentes]


def ler_planilhas(caminhos, mapa_territorios=None):
    """Lê as planilhas (uma por zip). Devolve {linhas: DataFrame, ressalvas, codigos_ausentes, nota_original}."""
    mapa = mapa_territorios or territorios()
    linhas, ressalvas, ausentes, notas_fonte = [], [], [], set()
    for caminho in caminhos:
        wb = openpyxl.load_workbook(caminho, data_only=True)
        etapas = set()
        for ws in wb.worksheets:
            l, r, a = ler_aba(ws, mapa, caminho.name)
            etapas.add(l[0]["etapa"] if l else None)
            linhas += l
            ressalvas += r
            ausentes += a
            for linha in range(1, ws.max_row + 1):
                texto = ws.cell(linha, 1).value
                if isinstance(texto, str) and texto.strip().startswith("Fonte:"):
                    notas_fonte.add(texto.strip())
        conferir(etapas == set(ETAPAS.values()), f"{caminho.name}: etapas {etapas}")
    df = pd.DataFrame(linhas)
    duplicadas = df.duplicated(["uf", "ano", "etapa", "rede", "indicador", "serie"], keep=False)
    conferir(not duplicadas.any(), f"linhas repetidas entre abas/arquivos: {df[duplicadas].head().to_dict('records')}")
    return {"linhas": df, "ressalvas": ressalvas, "codigos_ausentes": ausentes,
            "nota_original": " | ".join(sorted(notas_fonte)) or None}


# --- conferências -----------------------------------------------------------------------------------------

CONTROLES_2025 = [  # pedidos pelo usuário: (uf, etapa, rede, ideb publicado)
    ("RS", "anos_iniciais", "estadual", 6.4), ("RS", "anos_finais", "estadual", 5.4), ("RS", "medio", "estadual", 4.7),
    ("PR", "medio", "total", 5.1), ("PI", "medio", "total", 5.0), ("ES", "medio", "total", 4.9), ("RS", "medio", "total", 4.9),
]
TOLERANCIA_NXP = 0.1


def conferencias(df, ufs, anos_esperados):
    """Cobertura, N × P, controles 2025 e amplitude. Devolve um dict (o relatório)."""
    ideb = df[df["indicador"] == "ideb"]
    chave = ["uf", "etapa", "rede", "ano"]
    tabela = df[df["indicador"].isin(["ideb", "rendimento", "nota_padronizada"])].pivot_table(
        index=chave, columns="indicador", values="valor", aggfunc="first", dropna=False).reset_index()

    # Cobertura: por UF, etapa e rede, anos com IDEB publicado
    lacunas = []
    # Redes que a planilha traz para as UFs (a municipal só existe no arquivo do Brasil)
    redes_por_etapa = ideb[ideb["uf"].isin(ufs)].groupby("etapa")["rede"].unique().to_dict()
    for etapa, redes in sorted(redes_por_etapa.items()):
        for rede in sorted(redes):
            for uf in ufs:
                parte = ideb[(ideb["uf"] == uf) & (ideb["etapa"] == etapa) & (ideb["rede"] == rede)]
                faltando = sorted(set(anos_esperados) - set(parte.loc[parte["valor"].notna(), "ano"]))
                if faltando:
                    lacunas.append({"uf": uf, "etapa": etapa, "rede": rede, "anos_sem_ideb": faltando})
    cobertura = {
        "esperado": f"{len(ufs)} UFs × 3 etapas × {len(anos_esperados)} edições, por rede",
        "redes_por_etapa": {e: sorted(r) for e, r in redes_por_etapa.items()},
        "ideb_publicado_ufs": int(ideb[ideb["uf"].isin(ufs) & ideb["valor"].notna()].shape[0]),
        "lacunas": lacunas,
    }

    # IDEB = N × P (publicado arredondado a 1 casa)
    nxp = tabela.dropna(subset=["ideb", "rendimento", "nota_padronizada"]).copy()
    nxp["calculado"] = nxp["nota_padronizada"] * nxp["rendimento"]
    nxp["diferenca"] = (nxp["calculado"] - nxp["ideb"]).abs()
    fora = nxp[nxp["diferenca"] > TOLERANCIA_NXP + 1e-9]
    so_ideb = tabela[tabela["ideb"].notna() & (tabela["rendimento"].isna() | tabela["nota_padronizada"].isna())]

    controles = []
    for uf, etapa, rede, esperado in CONTROLES_2025:
        v = ideb[(ideb["uf"] == uf) & (ideb["etapa"] == etapa) & (ideb["rede"] == rede) & (ideb["ano"] == 2025)]["valor"]
        publicado = None if v.empty else v.iloc[0]
        controles.append({"uf": uf, "etapa": etapa, "rede": rede, "esperado": esperado, "publicado": publicado,
                          "ok": publicado is not None and abs(publicado - esperado) < 1e-9})

    def fora_da_faixa(indicador, minimo, maximo):
        parte = df[(df["indicador"] == indicador) & df["valor"].notna()]
        return parte[(parte["valor"] < minimo) | (parte["valor"] > maximo)][chave + ["valor"]].to_dict("records")

    return {
        "cobertura": cobertura,
        "n_x_p": {"tolerancia": TOLERANCIA_NXP, "conferidos": int(len(nxp)), "maior_diferenca": float(nxp["diferenca"].max()),
                  "fora_da_tolerancia": fora[chave + ["nota_padronizada", "rendimento", "calculado", "ideb", "diferenca"]].to_dict("records"),
                  "ideb_sem_n_ou_p": so_ideb[chave + ["ideb"]].to_dict("records")},
        "controles_2025": controles,
        "amplitude": {"ideb_fora_de_0_10": fora_da_faixa("ideb", 0, 10), "rendimento_fora_de_0_1": fora_da_faixa("rendimento", 0, 1),
                      "nota_padronizada_fora_de_0_10": fora_da_faixa("nota_padronizada", 0, 10),
                      "aprovacao_fora_de_0_100": fora_da_faixa("aprovacao", 0, 100)},
    }


def tabela(df, extraido_em, edicao):
    """Tabela `ideb` com cod_ibge dos dicionários (UFs, regiões e Brasil)."""
    codigos = pd.concat([pd.read_csv(c, dtype=str)[["uf", "cod_ibge"]]
                         for c in (base.CAMINHO_UFS, CAMINHO_REGIOES, CAMINHO_TERRITORIOS_NAO_UF)])
    saida = df.assign(cod_ibge=df["uf"].map(dict(zip(codigos["uf"], codigos["cod_ibge"]))))
    conferir(saida["cod_ibge"].notna().all(), f"sem código: {sorted(saida.loc[saida['cod_ibge'].isna(), 'uf'].unique())}")
    return saida.assign(edicao=edicao, fonte_dado=FONTE_DADO, extraido_em=extraido_em)


def main(argv=None):
    argv = sys.argv[1:] if argv is None else argv
    zips = baixar_mod.baixar(atualizar="--atualizar" in argv)
    manifesto = json.loads((baixar_mod.DIR_RAW / baixar_mod.MANIFESTO).read_text(encoding="utf-8"))
    arquivos, planilhas = [], []
    for arquivo, caminho_zip in zips.items():
        registro = manifesto[baixar_mod.relativo(arquivo)]
        conferir(hashlib.sha256(caminho_zip.read_bytes()).hexdigest() == registro["sha256"], f"{caminho_zip.name}: SHA-256 ≠ manifesto")
        pasta = caminho_zip.with_suffix("")
        [xlsx] = list(pasta.rglob("*.xlsx"))
        [md5_txt] = list(pasta.rglob("md5_*.txt"))
        md5_inep = next(l.split()[0].lower() for l in md5_txt.read_text(encoding="utf-8", errors="replace").splitlines() if xlsx.name in l)
        conferir(hashlib.md5(xlsx.read_bytes()).hexdigest() == md5_inep, f"{xlsx.name}: MD5 diferente do publicado pelo Inep")
        arquivos.append({"url": registro["url"], "arquivo": baixar_mod.relativo(arquivo), "sha256": registro["sha256"],
                         "planilha": xlsx.name, "md5_inep_conferido": md5_inep, "extraido_em": registro["extraido_em"]})
        planilhas.append(xlsx)

    lido = ler_planilhas(planilhas)
    df = lido["linhas"]
    ufs = pd.read_csv(base.CAMINHO_UFS, dtype=str)["uf"].tolist()
    anos = sorted(df.loc[df["indicador"] == "ideb", "ano"].unique().tolist())
    relatorio = conferencias(df, ufs, anos)

    # Valores de controle: se algum não bater, para antes de gravar qualquer coisa
    errados = [c for c in relatorio["controles_2025"] if not c["ok"]]
    if errados:
        raise base.ErroValidacao(f"valores de controle de 2025 não batem: {errados}")
    amplitude = {k: v for k, v in relatorio["amplitude"].items() if v}
    if amplitude:
        raise base.ErroValidacao(f"valores fora da faixa: {amplitude}")

    extraido_em = max(a["extraido_em"] for a in arquivos)
    gravados = base.gravar(tabela(df, extraido_em, baixar_mod.EDICAO), "ideb", FONTE_DADO)
    CAMINHO_FONTE.write_text(json.dumps({
        "edicao": baixar_mod.EDICAO, "arquivos": arquivos, "nota_original": lido["nota_original"], "anos": anos,
        "ressalvas_por_linha": lido["ressalvas"], "codigos_ausentes": lido["codigos_ausentes"], "conferencias": relatorio,
    }, ensure_ascii=False, indent=1, default=float), encoding="utf-8")

    fora = relatorio["n_x_p"]["fora_da_tolerancia"]
    print(f"{df['uf'].nunique()} territórios, anos {anos[0]}–{anos[-1]}, {len(df):,} valores; {len(gravados)} arquivos gravados")
    print(f"controles 2025: {sum(c['ok'] for c in relatorio['controles_2025'])}/{len(relatorio['controles_2025'])} ok; "
          f"N × P: {relatorio['n_x_p']['conferidos']} conferidos, {len(fora)} fora de ±{TOLERANCIA_NXP}; "
          f"lacunas (UF × etapa × rede com ano sem IDEB): {len(relatorio['cobertura']['lacunas'])}; "
          f"códigos ausentes: {len(lido['codigos_ausentes'])}")


if __name__ == "__main__":
    main()
